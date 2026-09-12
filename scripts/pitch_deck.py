#!/usr/bin/env python3
"""pitch_deck.py - render a deck spec into a .pptx you can open and edit.

The problem it exists for: a founder who needs an investor deck writes it in a
chat window, then retypes it into slides. The retyping is where the structure
gets lost and where the numbers drift from the ones that were checked. This
takes the markdown spec the `pitch-deck` skill produces and renders it, so the
spec stays the source and the slides are a render of it.

The spec is the deliverable. The .pptx is a convenience. That ordering is the
point: the markdown is yours, readable, diffable and editable with no software,
and if this script cannot run you still have the deck.

This is the SECOND script in the repo that is not standard library. It needs
`pip install python-pptx`, and like `scrape.py` it fails closed: without the
package it exits 1, prints the exact command, and names the fallback, so the
skill above it degrades instead of breaking.

One thing it must not do: credit itself. `deliverable_gate.py attribution`
fails any document whose author metadata names the library that generated it,
so the author is read from `os-config.yaml` and a deck with no configured
author is left unstamped rather than stamped "python-pptx".

Invariants: reads the spec and os-config.yaml, writes exactly one .pptx at the
path you name, never edits the spec, no network, no key, no model call.
Exit 0 written, 1 cannot run (missing package, missing spec, unwritable path).

Usage:
  python scripts/pitch_deck.py <spec.md>
  python scripts/pitch_deck.py <spec.md> --out decks/acme.pptx
  python scripts/pitch_deck.py <spec.md> --check      # parse only, write nothing
"""

from __future__ import annotations

import argparse
import datetime as _dt
import re
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]

PIP_LINE = (
    "pitch_deck.py needs one package: pip install python-pptx\n"
    "Without it the markdown spec is still your deck - open it, or paste it into "
    "Canva, Gamma, Pitch or Google Slides. Nothing is lost.\n"
)


def _force_utf8() -> None:
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass


# ------------------------------------------------------------------ spec --

class Slide:
    __slots__ = ("title", "bullets", "notes")

    def __init__(self, title: str):
        self.title = title
        self.bullets: list[str] = []
        self.notes: list[str] = []


class Spec:
    def __init__(self) -> None:
        self.title = ""
        self.subtitle = ""
        self.slides: list[Slide] = []


_BULLET = re.compile(r"^\s*[-*+]\s+(.*)$")
_NOTE = re.compile(r"^\s*>\s?(.*)$")
_META = re.compile(r"^\s*([A-Za-z_]+):\s*(.*)$")
_INLINE = [
    (re.compile(r"\[([^\]]+)\]\([^)]*\)"), r"\1"),
    (re.compile(r"`([^`]*)`"), r"\1"),
    (re.compile(r"\*\*|__"), ""),
]


def _clean(s: str) -> str:
    for pat, rep in _INLINE:
        s = pat.sub(rep, s)
    return s.strip()


def parse_spec(text: str) -> Spec:
    """Read the markdown spec. `#` is the deck, `##` is a slide, `-` a bullet,
    `>` a speaker note."""
    spec = Spec()
    current: Slide | None = None
    in_code = False
    for raw in text.splitlines():
        line = raw.rstrip()
        st = line.strip()
        if st.startswith("```"):
            in_code = not in_code
            continue
        if in_code or not st:
            continue
        if st.startswith("## "):
            current = Slide(_clean(st[3:]))
            spec.slides.append(current)
            continue
        if st.startswith("# "):
            spec.title = _clean(st[2:])
            current = None
            continue
        m = _NOTE.match(st)
        if m and current is not None:
            current.notes.append(_clean(m.group(1)))
            continue
        m = _BULLET.match(st)
        if m:
            if current is None:                     # a bullet before any slide
                continue
            current.bullets.append(_clean(m.group(1)))
            continue
        m = _META.match(st)
        if m and current is None and m.group(1).lower() == "subtitle":
            spec.subtitle = _clean(m.group(2))
            continue
        if current is not None:
            current.bullets.append(_clean(st))
    return spec


def read_brand(root: Path) -> dict:
    """Font and author from os-config.yaml. Unset stays unset: a wrong brand
    assumption is worse than an honest gap, which is the rule the ship gate
    already follows."""
    cfg = {"font": "", "document_author": ""}
    path = root / "os-config.yaml"
    if not path.is_file():
        return cfg
    section = None
    for raw in path.read_text(encoding="utf-8", errors="replace").splitlines():
        if raw.strip().startswith("#") or not raw.strip():
            continue
        if not raw.startswith((" ", "\t")):
            section = raw.split(":", 1)[0].strip()
            continue
        if section != "brand" or ":" not in raw:
            continue
        key, _, value = raw.strip().partition(":")
        if key.strip() in cfg:
            cfg[key.strip()] = value.strip().strip("\"'")
    return cfg


# ---------------------------------------------------------------- render --

def render(spec: Spec, out: Path, brand: dict) -> Path:
    try:
        from pptx import Presentation
        from pptx.util import Pt
    except ImportError:
        sys.stderr.write(PIP_LINE)
        raise SystemExit(1)

    prs = Presentation()
    title_layout = prs.slide_layouts[0]
    body_layout = prs.slide_layouts[1]

    if spec.title:
        s = prs.slides.add_slide(title_layout)
        s.shapes.title.text = spec.title
        if len(s.placeholders) > 1:
            s.placeholders[1].text = spec.subtitle or ""

    for slide in spec.slides:
        s = prs.slides.add_slide(body_layout)
        s.shapes.title.text = slide.title
        body = s.placeholders[1].text_frame
        body.clear()
        for i, bullet in enumerate(slide.bullets):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = bullet
            para.level = 0
        if slide.notes:
            s.notes_slide.notes_text_frame.text = "\n".join(slide.notes)

    font = brand.get("font") or ""
    if font:
        for s in prs.slides:
            for shape in s.shapes:
                if not shape.has_text_frame:
                    continue
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.name = font
                        if run.font.size is None:
                            run.font.size = Pt(18)

    # Never credit the library. deliverable_gate.py attribution fails a document
    # whose author names the tool that made it, and it should.
    author = brand.get("document_author") or ""
    props = prs.core_properties
    props.author = author
    props.last_modified_by = author
    props.title = spec.title or out.stem
    props.comments = ""
    # The blank template carries its own creation date from years ago, and a
    # deck whose properties say it was made in 2013 is a small lie that a
    # recipient can see in one click.
    now = _dt.datetime.now(_dt.timezone.utc).replace(tzinfo=None)
    props.created = now
    props.modified = now

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


def main(argv: list[str] | None = None) -> int:
    _force_utf8()
    ap = argparse.ArgumentParser(description="Render a deck spec into a .pptx.")
    ap.add_argument("spec", help="the markdown deck spec")
    ap.add_argument("--out", default=None, help="output .pptx (default: beside the spec)")
    ap.add_argument("--check", action="store_true", help="parse only, write nothing")
    ap.add_argument("--root", default=None, help="OS root override (for tests)")
    args = ap.parse_args(argv)

    spec_path = Path(args.spec)
    if not spec_path.is_file():
        print(f"pitch_deck: no spec at {spec_path}", file=sys.stderr)
        return 1

    spec = parse_spec(spec_path.read_text(encoding="utf-8", errors="replace"))
    if not spec.slides:
        print("pitch_deck: the spec has no '## ' slide headings - nothing to render",
              file=sys.stderr)
        return 1

    root = Path(args.root) if args.root else REPO_ROOT
    brand = read_brand(root)

    if args.check:
        print(f"pitch_deck: {spec_path.name} parses - "
              f"{len(spec.slides)} slides, "
              f"{sum(len(s.bullets) for s in spec.slides)} bullets, "
              f"{sum(1 for s in spec.slides if s.notes)} with speaker notes")
        if not brand.get("document_author"):
            print("  note: os-config.yaml has no brand.document_author, so the "
                  "rendered deck is left unstamped rather than credited to a tool")
        return 0

    out = Path(args.out) if args.out else spec_path.with_suffix(".pptx")
    written = render(spec, out, brand)
    print(f"pitch_deck: wrote {written} ({len(spec.slides)} slides)")
    if not brand.get("document_author"):
        print("  author metadata left empty - set brand.document_author in "
              "os-config.yaml and re-run to stamp it")
    return 0


if __name__ == "__main__":
    sys.exit(main())
